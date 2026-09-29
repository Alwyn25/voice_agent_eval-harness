"""
full_eval.py  —  Comprehensive end-to-end inbound-call evaluation.
Measures: wall_ms, TTIF, p50–p99, accuracy, tool-accuracy — broken down by
language, difficulty, slice, and archetype.

Multi-turn follow-up (v2):
  When tool_expected=True and the agent asks for name/date instead of calling
  the tool directly, the eval now sends a second turn ("Ravi Sharma, this
  Saturday at 11am") to simulate the caller providing the required info, then
  checks if the tool fires.  This gives a valid tool-accuracy baseline that
  accounts for the agent's correct multi-turn booking flow.

Hindi investigation:
  judge() now returns (score, reason).  The Accuracy_ST and Hindi_Analysis
  sheets break out accuracy, tool-accuracy, and judge_reason for each language
  so judge bias vs model language gap can be distinguished.

Usage:
    python eval_harness/full_eval.py --label inbound_full_v2
    python eval_harness/full_eval.py --label inbound_full_v2 --concurrency 1
"""
import argparse, asyncio, json, math, os, re, statistics, time
from collections import defaultdict
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path

import httpx
import pandas as pd

# ── Config ─────────────────────────────────────────────────────────────────────
BASE    = "http://localhost:8000"
DATASET = Path(r"C:\Users\ADVORA-ALWYN\OneDrive\Desktop\highland_greenz_golden_dataset.xlsx")
OUT_DIR = Path(__file__).parent / "results" / "full"
OUT_DIR.mkdir(parents=True, exist_ok=True)
EMAIL   = "admin@advora.ai"
PASS    = "Admin@123"

# Load .env if present
_env = Path(__file__).parent.parent / "AI-Voice-Agent-Backend-main" / ".env"
if _env.exists():
    for _ln in _env.read_text(encoding="utf-8").splitlines():
        _ln = _ln.strip()
        if _ln and not _ln.startswith("#") and "=" in _ln:
            k, _, v = _ln.partition("=")
            if k.strip() not in os.environ:
                os.environ[k.strip()] = v.strip()

# ── Regex patterns ──────────────────────────────────────────────────────────────
CONFIRM_RE = re.compile(
    r"(shall i|should i|would you like me to|want me to|go ahead|confirm|proceed|"
    r"is that right|can i go ahead|shall we|do you want me)",
    re.IGNORECASE,
)
# Detects when the agent is collecting required info before booking
CLARIFY_NAME_RE = re.compile(
    r"(your name|may i (know|have|get) (your )?name|could i (have|get) (your )?name|"
    r"naam|aapka naam|what('s| is) your name)",
    re.IGNORECASE,
)
CLARIFY_DATE_RE = re.compile(
    r"(which date|what date|preferred date|date works|when (would|are) you|"
    r"kab|konsi date|what time|preferred time)",
    re.IGNORECASE,
)

_JUDGE_URL = "https://api.openai.com/v1/chat/completions"

# ── Helpers ─────────────────────────────────────────────────────────────────────
class TokenKeeper:
    def __init__(self): self._token = ""; self._issued = 0.0

    async def fresh(self, client: httpx.AsyncClient) -> str:
        if time.monotonic() - self._issued > 1400:
            r = await client.post(f"{BASE}/auth/login",
                                  json={"email": EMAIL, "password": PASS}, timeout=15)
            r.raise_for_status()
            self._token = r.json()["access_token"]
            self._issued = time.monotonic()
        return self._token

async def new_session(client, token, direction="inbound"):
    r = await client.post(f"{BASE}/sessions",
                          json={"channel": "phone", "direction": direction,
                                "purpose": "sales", "enforce_outbound_policy": False},
                          headers={"Authorization": f"Bearer {token}"}, timeout=20)
    r.raise_for_status()
    return r.json()["session_id"]

async def send_turn(client, token, sid, text):
    t0 = time.perf_counter()
    r  = await client.post(f"{BASE}/sessions/{sid}/turns",
                           json={"text": text},
                           headers={"Authorization": f"Bearer {token}"}, timeout=120)
    wall_ms = int((time.perf_counter() - t0) * 1000)
    r.raise_for_status()
    return r.json(), wall_ms

async def end_session(client, token, sid):
    try:
        await client.post(f"{BASE}/sessions/{sid}/end",
                          json={"analyse": "skip"},
                          headers={"Authorization": f"Bearer {token}"}, timeout=10)
    except Exception:
        pass

def parse_response(data, wall_ms):
    debug    = data.get("debug") or {}
    timing   = debug.get("timing") or {}
    executed = data.get("executed_tools") or []
    if isinstance(executed, list):
        tool_names = [(t.get("tool") or t.get("name") or str(t))
                      if isinstance(t, dict) else str(t) for t in executed]
    else:
        tool_names = [str(executed)] if executed else []
    pm   = timing.get("planner_ms", 0) or 0
    rm   = timing.get("retrieval_ms", 0) or 0
    tm   = timing.get("tool_ms", 0) or 0
    ttif = pm + rm + tm
    return {
        "response_text":    data.get("response_text", ""),
        "wall_ms":          timing.get("wall_ms", 0) or wall_ms,
        "planner_ms":       pm,
        "retrieval_ms":     rm,
        "tool_ms":          tm,
        "responder_ms":     timing.get("responder_ms", 0) or 0,
        "ttif_ms":          ttif,
        "fast_path":        (pm == 0 and ttif < 20),
        "tool_called":      len(tool_names) > 0,
        "executed_tools":   ", ".join(tool_names),
        "responder_source": timing.get("responder_source", "llm"),
    }

async def judge(client, utterance, response, required, forbidden):
    """Returns (score: float, reason: str).  score=-1 if judge not available."""
    key = os.environ.get("CI_LLM_API_KEY") or os.environ.get("OPENAI_API_KEY", "")
    if not key or (not str(required or "").strip() and not str(forbidden or "").strip()):
        return -1.0, ""
    prompt = (
        f"Score the agent response 0-100.\n"
        f"Required facts (ALL must appear): {required or 'none'}\n"
        f"Forbidden facts (must NOT appear): {forbidden or 'none'}\n"
        f"Customer: {utterance}\nAgent: {response}\n"
        f'JSON only: {{"score":<int>,"reason":"<12 words>"}}'
    )
    try:
        r = await client.post(_JUDGE_URL,
            json={"model": "gpt-4o-mini", "temperature": 0, "max_tokens": 80,
                  "messages": [{"role": "user", "content": prompt}]},
            headers={"Authorization": f"Bearer {key}"}, timeout=30)
        r.raise_for_status()
        obj = json.loads(r.json()["choices"][0]["message"]["content"])
        return float(obj.get("score", -1)), str(obj.get("reason", ""))
    except Exception:
        return -1.0, ""

# ── Row dataclass ──────────────────────────────────────────────────────────────
@dataclass
class Row:
    test_type:        str   = ""
    label:            str   = ""
    row_id:           str   = ""
    thread_id:        str   = ""
    turn_no:          int   = 1
    lang:             str   = ""
    difficulty:       str   = ""
    slice_type:       str   = ""    # rag / tool / escalation / memory / smalltalk …
    expected_intent:  str   = ""    # I01 … I29
    archetype:        str   = ""    # thread archetype (price_to_booking_books …)
    utterance:        str   = ""
    tool_expected:    bool  = False
    wall_ms:          int   = 0
    planner_ms:       int   = 0
    retrieval_ms:     int   = 0
    tool_ms:          int   = 0
    responder_ms:     int   = 0
    ttif_ms:          int   = 0
    fast_path:        bool  = False
    tool_called:      bool  = False
    tool_match:       bool  = False
    tool_followup:    bool  = False  # True when tool fired only after a follow-up clarification turn
    accuracy_score:   float = -1.0
    judge_reason:     str   = ""
    responder_source: str   = "llm"
    response_text:    str   = ""
    executed_tools:   str   = ""
    direction:        str   = "inbound"
    path:             str   = ""
    timestamp:        str   = ""
    error:            str   = ""


def _now() -> str:
    return datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S")


def _parse_exp_tools(raw) -> list[str]:
    s = str(raw or "").strip()
    return [x.strip() for x in s.replace("[","").replace("]","").split(",")
            if x.strip() and x.strip() not in ("None","nan","")]


# ── Single-turn ────────────────────────────────────────────────────────────────
async def run_st(sem, client, meta, keeper, idx, total, label):
    async with sem:
        token    = await keeper.fresh(client)
        utt      = str(meta.get("user_utterance", ""))
        exp_list = _parse_exp_tools(meta.get("expected_tool_calls", ""))
        row = Row(
            test_type="st", label=label,
            row_id=str(meta.get("id","")),
            lang=str(meta.get("lang","")),
            difficulty=str(meta.get("difficulty","")),
            slice_type=str(meta.get("slice","")),
            expected_intent=str(meta.get("expected_intent","")),
            utterance=utt,
            tool_expected=bool(exp_list),
            direction=str(meta.get("channel","inbound")),
        )
        print(f"  ST [{idx:>3}/{total}] {row.row_id:<8} {row.lang:<9} {row.difficulty:<6}",
              end=" ", flush=True)
        try:
            sid  = await new_session(client, token)
            data, wall_ms = await send_turn(client, token, sid, utt)
            p    = parse_response(data, wall_ms)
            row.wall_ms = p["wall_ms"]; row.planner_ms = p["planner_ms"]
            row.retrieval_ms = p["retrieval_ms"]; row.tool_ms = p["tool_ms"]
            row.responder_ms = p["responder_ms"]; row.ttif_ms = p["ttif_ms"]
            row.fast_path = p["fast_path"]; row.tool_called = p["tool_called"]
            row.responder_source = p["responder_source"]
            row.response_text  = p["response_text"]
            row.executed_tools = p["executed_tools"]
            row.path      = ("fast_path" if row.fast_path
                             else ("tool" if row.tool_called else "planner"))
            row.timestamp = _now()

            # ── Confirm flow (agent asks "shall I…") ──────────────────────────
            if not row.tool_called and CONFIRM_RE.search(p["response_text"]):
                token = await keeper.fresh(client)
                d2, _ = await send_turn(client, token, sid, "Yes, go ahead please.")
                p2 = parse_response(d2, 0)
                if p2["tool_called"]:
                    row.tool_called = True; row.executed_tools = p2["executed_tools"]
                    row.path = "tool"

            # ── Multi-turn follow-up: agent asks for name/date (v2) ───────────
            if row.tool_expected and not row.tool_called:
                resp_txt = p["response_text"]
                needs_name = bool(CLARIFY_NAME_RE.search(resp_txt))
                needs_date = bool(CLARIFY_DATE_RE.search(resp_txt))
                if needs_name or needs_date:
                    followup = "Ravi Sharma, this Saturday at 11am please."
                    token = await keeper.fresh(client)
                    d3, _ = await send_turn(client, token, sid, followup)
                    p3 = parse_response(d3, 0)
                    if p3["tool_called"]:
                        row.tool_called = True; row.executed_tools = p3["executed_tools"]
                        row.path = "tool"; row.tool_followup = True

            row.tool_match = (row.tool_called == row.tool_expected)
            row.accuracy_score, row.judge_reason = await judge(
                client, utt, p["response_text"],
                meta.get("required_facts",""), meta.get("forbidden_facts",""))
            await end_session(client, token, sid)
        except Exception as e:
            row.error = str(e)[:200]
            if not row.timestamp: row.timestamp = _now()

        flag = "*" if row.fast_path else ("+" if row.tool_match else " ")
        fu   = " [fu]" if row.tool_followup else ""
        print(f"{flag}{fu} acc={row.accuracy_score:>3.0f} wall={row.wall_ms:>5}ms ttif={row.ttif_ms:>5}ms")
        return row


# ── Thread ─────────────────────────────────────────────────────────────────────
async def run_thread(client, thread_rows, keeper, t_idx, n_threads, label):
    token = await keeper.fresh(client)
    tid       = thread_rows[0].get("thread_id","")
    lang      = str(thread_rows[0].get("lang",""))
    archetype = str(thread_rows[0].get("archetype",""))
    print(f"  TH [{t_idx:>2}/{n_threads}] {tid} ({lang}) [{archetype}]")
    results = []
    try:
        sid = await new_session(client, token)
        for row_meta in thread_rows:
            token    = await keeper.fresh(client)
            utt      = str(row_meta.get("utterance",""))
            exp_list = _parse_exp_tools(row_meta.get("expected_tool_calls",""))
            row = Row(
                test_type="th", label=label, thread_id=tid,
                turn_no=int(row_meta.get("turn_no",1)),
                lang=lang, archetype=archetype,
                expected_intent=str(row_meta.get("expected_intent","")),
                utterance=utt, tool_expected=bool(exp_list),
                direction=str(row_meta.get("channel","inbound")),
            )
            try:
                data, wall_ms = await send_turn(client, token, sid, utt)
                p = parse_response(data, wall_ms)
                row.wall_ms = p["wall_ms"]; row.planner_ms = p["planner_ms"]
                row.retrieval_ms = p["retrieval_ms"]; row.tool_ms = p["tool_ms"]
                row.responder_ms = p["responder_ms"]; row.ttif_ms = p["ttif_ms"]
                row.fast_path = p["fast_path"]; row.tool_called = p["tool_called"]
                row.responder_source = p["responder_source"]
                row.response_text  = p["response_text"]
                row.executed_tools = p["executed_tools"]
                row.path      = ("fast_path" if row.fast_path
                                 else ("tool" if row.tool_called else "planner"))
                row.timestamp = _now()

                if not row.tool_called and CONFIRM_RE.search(p["response_text"]):
                    token = await keeper.fresh(client)
                    d2, _ = await send_turn(client, token, sid, "Yes, go ahead please.")
                    p2 = parse_response(d2, 0)
                    if p2["tool_called"]:
                        row.tool_called = True; row.executed_tools = p2["executed_tools"]
                        row.path = "tool"

                if row.tool_expected and not row.tool_called:
                    resp_txt = p["response_text"]
                    if CLARIFY_NAME_RE.search(resp_txt) or CLARIFY_DATE_RE.search(resp_txt):
                        followup = "Ravi Sharma, this Saturday at 11am please."
                        token = await keeper.fresh(client)
                        d3, _ = await send_turn(client, token, sid, followup)
                        p3 = parse_response(d3, 0)
                        if p3["tool_called"]:
                            row.tool_called = True; row.executed_tools = p3["executed_tools"]
                            row.path = "tool"; row.tool_followup = True

                row.tool_match = (row.tool_called == row.tool_expected)
                row.accuracy_score, row.judge_reason = await judge(
                    client, utt, p["response_text"],
                    row_meta.get("required_facts",""), row_meta.get("forbidden_facts",""))
            except Exception as e:
                row.error = str(e)[:200]
                if not row.timestamp: row.timestamp = _now()

            flag = "*" if row.fast_path else ("+" if row.tool_match else " ")
            fu   = "[fu]" if row.tool_followup else ""
            print(f"      {flag}{fu} turn {row.turn_no}  "
                  f"wall={row.wall_ms:>5}ms  ttif={row.ttif_ms:>5}ms  acc={row.accuracy_score:>3.0f}")
            results.append(row)
        await end_session(client, token, sid)
    except Exception as e:
        print(f"    THREAD FAILED: {e}")
    return results


# ── Statistics ─────────────────────────────────────────────────────────────────
def pct(vals, p):
    if not vals: return 0
    s = sorted(vals)
    idx = min(max(int(math.ceil(p / 100 * len(s))) - 1, 0), len(s) - 1)
    return s[idx]

def compute_stats(rows, label=""):
    ok       = [r for r in rows if not r.error and r.wall_ms > 0]
    walls    = [r.wall_ms   for r in ok]
    ttifs    = [r.ttif_ms   for r in ok if not r.fast_path]
    accs     = [r.accuracy_score for r in rows if r.accuracy_score >= 0]
    fp_rows  = [r for r in ok if r.fast_path]
    tool_exp = [r for r in rows if r.tool_expected]
    tool_hit = [r for r in rows if r.tool_expected and r.tool_match]
    errors   = sum(1 for r in rows if r.error)
    return {
        "label":         label,
        "n":             len(rows),
        "errors":        errors,
        "wall_mean":     round(statistics.mean(walls)) if walls else 0,
        "wall_p50":      pct(walls, 50),
        "wall_p75":      pct(walls, 75),
        "wall_p90":      pct(walls, 90),
        "wall_p95":      pct(walls, 95),
        "wall_p99":      pct(walls, 99),
        "ttif_mean":     round(statistics.mean(ttifs)) if ttifs else 0,
        "ttif_p50":      pct(ttifs, 50),
        "ttif_p75":      pct(ttifs, 75),
        "ttif_p90":      pct(ttifs, 90),
        "ttif_p95":      pct(ttifs, 95),
        "acc_mean":      round(statistics.mean(accs), 1) if accs else -1,
        "fast_path_n":   len(fp_rows),
        "tool_n":        len(tool_exp),
        "tool_accuracy": round(100 * len(tool_hit) / len(tool_exp), 1) if tool_exp else None,
    }

def print_summary(all_rows, label):
    W = 80
    print("\n" + "=" * W)
    print(f"  FULL EVAL — {label.upper()}")
    print("=" * W)
    segments = [
        ("Single-Turn", [r for r in all_rows if r.test_type == "st"]),
        ("Thread",      [r for r in all_rows if r.test_type == "th"]),
        ("OVERALL",     all_rows),
    ]
    for seg_name, rows in segments:
        if not rows: continue
        s = compute_stats(rows)
        print(f"\n  -- {seg_name} (n={s['n']}) --")
        print(f"  wall: p50={s['wall_p50']}ms  p95={s['wall_p95']}ms  mean={s['wall_mean']}ms")
        print(f"  TTIF: p50={s['ttif_p50']}ms  p95={s['ttif_p95']}ms")
        print(f"  acc={s['acc_mean']}  tool_acc={s['tool_accuracy']}%  fp={s['fast_path_n']}  err={s['errors']}")
    for lang in ("en","hi","hinglish"):
        rows = [r for r in all_rows if r.test_type == "st" and r.lang == lang]
        if not rows: continue
        s = compute_stats(rows, lang)
        print(f"  ST {lang:<10} acc={s['acc_mean']:>5}  tool_acc={s['tool_accuracy']}%")
    print("\n" + "=" * W)


# ── Excel output ──────────────────────────────────────────────────────────────
def _next_run_number(directory: Path) -> int:
    nums = []
    for p in directory.glob("eval_???.xlsx"):
        try: nums.append(int(p.name[5:8]))
        except ValueError: pass
    return max(nums, default=0) + 1


def save_results(all_rows, label):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    # ── JSONL ─────────────────────────────────────────────────────────────────
    jsonl_path = OUT_DIR / f"full_{label}.jsonl"
    with open(jsonl_path, "w", encoding="utf-8") as f:
        for r in all_rows:
            f.write(json.dumps(asdict(r), ensure_ascii=False) + "\n")

    run_no  = _next_run_number(OUT_DIR)
    xl_path = OUT_DIR / f"eval_{run_no:03d}_{label}.xlsx"
    wb      = openpyxl.Workbook()

    # ── Styles ─────────────────────────────────────────────────────────────────
    HDR_FILL  = PatternFill("solid", fgColor="1F3864")
    HDR_FONT  = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
    SUB_FILL  = PatternFill("solid", fgColor="2E75B6")
    SUB_FONT  = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
    SEC_FILL  = PatternFill("solid", fgColor="D9E2F3")
    SEC_FONT  = Font(name="Calibri", bold=True, color="1F3864", size=11)
    OK_FILL   = PatternFill("solid", fgColor="E2EFDA")
    WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
    BAD_FILL  = PatternFill("solid", fgColor="FCE4D6")
    ALT_FILL  = PatternFill("solid", fgColor="EBF3FB")
    CELL_FONT = Font(name="Calibri", size=10)
    _thin     = Side(style="thin", color="D0D0D0")
    BORDER    = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

    def hdr(ws, titles):
        ws.append(titles)
        for cell in ws[ws.max_row]:
            cell.fill = HDR_FILL; cell.font = HDR_FONT; cell.border = BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[ws.max_row].height = 28

    def sub_hdr(ws, titles):
        ws.append(titles)
        for cell in ws[ws.max_row]:
            cell.fill = SUB_FILL; cell.font = SUB_FONT; cell.border = BORDER
            cell.alignment = Alignment(horizontal="center")

    def sec_title(ws, title, n_cols=16):
        ws.append([title])
        end_col = get_column_letter(n_cols)
        ws.merge_cells(f"A{ws.max_row}:{end_col}{ws.max_row}")
        ws[f"A{ws.max_row}"].font = SEC_FONT
        ws[f"A{ws.max_row}"].fill = SEC_FILL

    def data_row(ws, vals, acc_col=None, err_col=None):
        ws.append(vals)
        row_idx = ws.max_row
        for j, cell in enumerate(ws[row_idx], 1):
            cell.font = CELL_FONT; cell.border = BORDER
            v = cell.value
            if acc_col and j == acc_col and isinstance(v, (int, float)) and v >= 0:
                cell.fill = OK_FILL if v >= 70 else WARN_FILL if v >= 40 else BAD_FILL
            elif err_col and j == err_col and v:
                cell.fill = BAD_FILL

    # ── Sheet 1: Summary ───────────────────────────────────────────────────────
    ws_sum = wb.active
    ws_sum.title = "Summary"
    ws_sum.append([f"EVAL SUMMARY  |  Run #{run_no}  |  Label: {label}",
                   "", "", f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M')} UTC"])
    ws_sum.merge_cells("A1:D1")
    ws_sum["A1"].font = Font(name="Calibri", bold=True, size=14, color="1F3864")
    ws_sum["A1"].fill = SEC_FILL
    ws_sum.row_dimensions[1].height = 26
    ws_sum.append([])

    STAT_HDR = ["Segment", "n", "Errors", "wall_mean", "wall_p50", "wall_p75",
                "wall_p90", "wall_p95", "wall_p99", "ttif_mean", "ttif_p50",
                "ttif_p90", "ttif_p95", "acc_mean", "fast_path_n", "tool_n", "tool_acc_%"]

    def _stat_row(ws, rows, seg_name):
        s  = compute_stats(rows, seg_name)
        ta = s["tool_accuracy"] if s["tool_accuracy"] is not None else ""
        data_row(ws, [seg_name, s["n"], s["errors"],
                      s["wall_mean"], s["wall_p50"], s["wall_p75"],
                      s["wall_p90"], s["wall_p95"], s["wall_p99"],
                      s["ttif_mean"], s["ttif_p50"], s["ttif_p90"], s["ttif_p95"],
                      s["acc_mean"], s["fast_path_n"], s["tool_n"], ta])

    sub_hdr(ws_sum, STAT_HDR)
    for seg_name, rows in [
        ("Single-Turn",  [r for r in all_rows if r.test_type == "st"]),
        ("Threads",      [r for r in all_rows if r.test_type == "th"]),
        ("OVERALL",      all_rows),
    ]:
        if rows: _stat_row(ws_sum, rows, seg_name)
    ws_sum.append([])

    for title, grp_keys, type_f, attr in [
        ("By Language — Single-Turn",   ["en","hi","hinglish"],     "st", "lang"),
        ("By Difficulty — Single-Turn", ["easy","medium","hard"],   "st", "difficulty"),
        ("By Slice — Single-Turn",      ["rag","tool","escalation",
                                         "memory","smalltalk","voice",
                                         "adversarial","qualification",
                                         "objection","compliance"],  "st", "slice_type"),
        ("By Language — Threads",       ["en","hi","hinglish"],     "th", "lang"),
        ("By Archetype — Threads",      None,                       "th", "archetype"),
    ]:
        sec_title(ws_sum, title, len(STAT_HDR))
        sub_hdr(ws_sum, STAT_HDR)
        keys = grp_keys or sorted({getattr(r, attr, "") for r in all_rows
                                    if r.test_type == type_f and getattr(r, attr, "")})
        for key in keys:
            rows = [r for r in all_rows
                    if r.test_type == type_f and getattr(r, attr, "") == key]
            if rows: _stat_row(ws_sum, rows, key)
        ws_sum.append([])

    for ci, w in enumerate([28,5,6,10,9,9,9,9,9,10,9,9,9,9,11,7,9], 1):
        ws_sum.column_dimensions[get_column_letter(ci)].width = w
    ws_sum.freeze_panes = "A3"

    # ── Sheet 2: Accuracy_ST ──────────────────────────────────────────────────
    ws_acc_st = wb.create_sheet("Accuracy_ST")
    sec_title(ws_acc_st, "Single-Turn Accuracy Matrix", 8)
    ws_acc_st.append([])

    # Grid: Language × Difficulty
    sec_title(ws_acc_st, "Accuracy by Language × Difficulty (mean score)", 8)
    diffs = ["easy", "medium", "hard", "ALL"]
    sub_hdr(ws_acc_st, ["Lang \\ Diff"] + diffs + ["tool_acc_%"])
    for lang in ("en", "hi", "hinglish"):
        rowvals = [lang]
        for diff in diffs:
            filt = [r for r in all_rows if r.test_type == "st" and r.lang == lang
                    and (diff == "ALL" or r.difficulty == diff)]
            accs = [r.accuracy_score for r in filt if r.accuracy_score >= 0]
            rowvals.append(round(statistics.mean(accs), 1) if accs else "")
        tool_r = [r for r in all_rows if r.test_type == "st" and r.lang == lang]
        te = [r for r in tool_r if r.tool_expected]
        th_ = [r for r in tool_r if r.tool_expected and r.tool_match]
        rowvals.append(round(100*len(th_)/len(te), 1) if te else "")
        ws_acc_st.append(rowvals)
        for cell in ws_acc_st[ws_acc_st.max_row]:
            cell.font = CELL_FONT; cell.border = BORDER
    ws_acc_st.append([])

    # Grid: Slice breakdown
    sec_title(ws_acc_st, "Accuracy by Slice × Language", 6)
    sub_hdr(ws_acc_st, ["Slice", "n", "acc_en", "acc_hi", "acc_hinglish", "acc_ALL", "tool_acc_%"])
    slices = ["rag","tool","escalation","memory","smalltalk","voice",
              "adversarial","qualification","objection","compliance"]
    for sl in slices:
        rowvals = [sl]
        sl_all = [r for r in all_rows if r.test_type == "st" and r.slice_type == sl]
        rowvals.append(len(sl_all))
        for lang in ("en","hi","hinglish"):
            filt = [r for r in sl_all if r.lang == lang]
            accs = [r.accuracy_score for r in filt if r.accuracy_score >= 0]
            rowvals.append(round(statistics.mean(accs), 1) if accs else "")
        accs_all = [r.accuracy_score for r in sl_all if r.accuracy_score >= 0]
        rowvals.append(round(statistics.mean(accs_all), 1) if accs_all else "")
        te = [r for r in sl_all if r.tool_expected]
        th_ = [r for r in sl_all if r.tool_expected and r.tool_match]
        rowvals.append(round(100*len(th_)/len(te), 1) if te else "")
        ws_acc_st.append(rowvals)
        for cell in ws_acc_st[ws_acc_st.max_row]:
            cell.font = CELL_FONT; cell.border = BORDER
    ws_acc_st.append([])

    # Tool accuracy with vs without follow-up
    sec_title(ws_acc_st, "Tool Accuracy — Single-Turn (with multi-turn follow-up v2)", 6)
    sub_hdr(ws_acc_st, ["Lang", "tool_expected_n", "fired_t1_only",
                         "fired_via_followup", "total_matched", "tool_acc_%"])
    for lang in ("en","hi","hinglish","ALL"):
        lang_rows = [r for r in all_rows
                     if r.test_type == "st" and (lang == "ALL" or r.lang == lang)]
        te   = [r for r in lang_rows if r.tool_expected]
        hit  = [r for r in te if r.tool_match]
        t1   = [r for r in hit if not r.tool_followup]
        fu   = [r for r in hit if r.tool_followup]
        ws_acc_st.append([lang, len(te), len(t1), len(fu), len(hit),
                          round(100*len(hit)/len(te), 1) if te else ""])
        for cell in ws_acc_st[ws_acc_st.max_row]:
            cell.font = CELL_FONT; cell.border = BORDER
    ws_acc_st.append([])

    for ci, w in enumerate([20,8,9,9,12,10,10], 1):
        ws_acc_st.column_dimensions[get_column_letter(ci)].width = w
    ws_acc_st.freeze_panes = "A2"

    # ── Sheet 3: Accuracy_TH ──────────────────────────────────────────────────
    ws_acc_th = wb.create_sheet("Accuracy_TH")
    sec_title(ws_acc_th, "Thread Accuracy Matrix", 8)
    ws_acc_th.append([])

    # Grid: Archetype × Language
    sec_title(ws_acc_th, "Accuracy by Archetype × Language", 8)
    archetypes = sorted({r.archetype for r in all_rows if r.test_type == "th" and r.archetype})
    sub_hdr(ws_acc_th, ["Archetype", "n_turns", "acc_en", "acc_hi", "acc_hinglish",
                          "acc_ALL", "tool_acc_%", "followup_n"])
    for arch in archetypes:
        arch_rows = [r for r in all_rows if r.test_type == "th" and r.archetype == arch]
        rowvals = [arch, len(arch_rows)]
        for lang in ("en","hi","hinglish"):
            filt = [r for r in arch_rows if r.lang == lang]
            accs = [r.accuracy_score for r in filt if r.accuracy_score >= 0]
            rowvals.append(round(statistics.mean(accs), 1) if accs else "")
        accs_all = [r.accuracy_score for r in arch_rows if r.accuracy_score >= 0]
        rowvals.append(round(statistics.mean(accs_all), 1) if accs_all else "")
        te  = [r for r in arch_rows if r.tool_expected]
        th_ = [r for r in arch_rows if r.tool_expected and r.tool_match]
        rowvals.append(round(100*len(th_)/len(te), 1) if te else "")
        rowvals.append(sum(1 for r in arch_rows if r.tool_followup))
        ws_acc_th.append(rowvals)
        for cell in ws_acc_th[ws_acc_th.max_row]:
            cell.font = CELL_FONT; cell.border = BORDER
    ws_acc_th.append([])

    # By turn number
    sec_title(ws_acc_th, "Accuracy by Turn Number within Thread", 6)
    sub_hdr(ws_acc_th, ["Turn #", "n", "acc_en", "acc_hi", "acc_hinglish", "acc_ALL", "tool_acc_%"])
    max_turn = max((r.turn_no for r in all_rows if r.test_type == "th"), default=1)
    for tn in range(1, max_turn + 1):
        turn_rows = [r for r in all_rows if r.test_type == "th" and r.turn_no == tn]
        if not turn_rows: continue
        rowvals = [tn, len(turn_rows)]
        for lang in ("en","hi","hinglish"):
            accs = [r.accuracy_score for r in turn_rows
                    if r.lang == lang and r.accuracy_score >= 0]
            rowvals.append(round(statistics.mean(accs), 1) if accs else "")
        accs_all = [r.accuracy_score for r in turn_rows if r.accuracy_score >= 0]
        rowvals.append(round(statistics.mean(accs_all), 1) if accs_all else "")
        te  = [r for r in turn_rows if r.tool_expected]
        th_ = [r for r in turn_rows if r.tool_expected and r.tool_match]
        rowvals.append(round(100*len(th_)/len(te), 1) if te else "")
        ws_acc_th.append(rowvals)
        for cell in ws_acc_th[ws_acc_th.max_row]:
            cell.font = CELL_FONT; cell.border = BORDER

    for ci, w in enumerate([32,8,9,9,13,9,10,11], 1):
        ws_acc_th.column_dimensions[get_column_letter(ci)].width = w
    ws_acc_th.freeze_panes = "A2"

    # ── Sheet 4: Hindi_Analysis ───────────────────────────────────────────────
    ws_hi = wb.create_sheet("Hindi_Analysis")
    sec_title(ws_hi, "Hindi Gap Investigation — Why do Hindi queries score ~8pp lower?", 10)
    ws_hi.append([])

    # Per-slice gap: Hindi vs English
    sec_title(ws_hi, "Accuracy Gap by Slice (EN − HI)", 5)
    sub_hdr(ws_hi, ["Slice", "n_en", "acc_en", "n_hi", "acc_hi", "gap (en-hi)",
                     "n_hinglish", "acc_hinglish"])
    for sl in slices:
        en_r = [r for r in all_rows if r.test_type == "st" and r.slice_type == sl and r.lang == "en"]
        hi_r = [r for r in all_rows if r.test_type == "st" and r.slice_type == sl and r.lang == "hi"]
        hg_r = [r for r in all_rows if r.test_type == "st" and r.slice_type == sl and r.lang == "hinglish"]
        en_a = round(statistics.mean([r.accuracy_score for r in en_r if r.accuracy_score >= 0]), 1) if en_r else ""
        hi_a = round(statistics.mean([r.accuracy_score for r in hi_r if r.accuracy_score >= 0]), 1) if hi_r else ""
        hg_a = round(statistics.mean([r.accuracy_score for r in hg_r if r.accuracy_score >= 0]), 1) if hg_r else ""
        gap  = round(float(en_a or 0) - float(hi_a or 0), 1) if en_a != "" and hi_a != "" else ""
        ws_hi.append([sl, len(en_r), en_a, len(hi_r), hi_a, gap, len(hg_r), hg_a])
        row_obj = ws_hi[ws_hi.max_row]
        for cell in row_obj: cell.font = CELL_FONT; cell.border = BORDER
        if isinstance(gap, (int, float)):
            row_obj[5].fill = BAD_FILL if gap > 10 else WARN_FILL if gap > 5 else OK_FILL
    ws_hi.append([])

    # Low-scoring Hindi rows: score < 40
    sec_title(ws_hi, "Hindi rows with accuracy < 40 — judge reasons", 8)
    sub_hdr(ws_hi, ["row_id", "slice", "intent", "difficulty", "utterance (Hindi)",
                     "response_text", "required_facts", "accuracy_score", "judge_reason"])
    lo_hi = sorted(
        [r for r in all_rows if r.test_type == "st" and r.lang == "hi"
         and 0 <= r.accuracy_score < 40],
        key=lambda r: r.accuracy_score)
    # We don't store required_facts in Row, so just show what we have
    for r in lo_hi:
        ws_hi.append([r.row_id, r.slice_type, r.expected_intent, r.difficulty,
                      r.utterance, r.response_text[:200],
                      "", r.accuracy_score, r.judge_reason])
        for cell in ws_hi[ws_hi.max_row]: cell.font = CELL_FONT; cell.border = BORDER
    ws_hi.append([])

    # Judge bias test: compare judge_reason text patterns for EN vs HI
    sec_title(ws_hi, "Judge Reason Pattern Analysis (EN vs HI low-scorers < 40)", 4)
    sub_hdr(ws_hi, ["Lang", "n_low (<40)", "% of lang total",
                     "common reason keywords"])
    for lang in ("en","hi","hinglish"):
        all_lang = [r for r in all_rows if r.test_type == "st" and r.lang == lang]
        lo = [r for r in all_lang if 0 <= r.accuracy_score < 40]
        reasons = " ".join(r.judge_reason.lower() for r in lo if r.judge_reason)
        # Count top keyword patterns
        kw_counts = defaultdict(int)
        for kw in ["missing", "not mentioned", "incorrect", "wrong", "language",
                   "hindi", "translation", "fact", "price", "required"]:
            kw_counts[kw] = reasons.count(kw)
        top_kw = ", ".join(f"{k}×{v}" for k, v in
                           sorted(kw_counts.items(), key=lambda x: -x[1]) if v > 0)[:120]
        pct_lo = round(100 * len(lo) / len(all_lang), 1) if all_lang else ""
        ws_hi.append([lang, len(lo), pct_lo, top_kw])
        for cell in ws_hi[ws_hi.max_row]: cell.font = CELL_FONT; cell.border = BORDER

    for ci, w in enumerate([14,12,10,10,45,55,45,14,50], 1):
        ws_hi.column_dimensions[get_column_letter(ci)].width = w
    ws_hi.freeze_panes = "A2"

    # ── Sheet 5: Raw_ST ───────────────────────────────────────────────────────
    ws_st = wb.create_sheet("Raw_ST")
    RAW_COLS = [
        ("Run #",          None,               7),
        ("Label",          "label",            14),
        ("Row ID",         "row_id",           10),
        ("Lang",           "lang",             9),
        ("Difficulty",     "difficulty",       10),
        ("Slice",          "slice_type",       12),
        ("Intent",         "expected_intent",  8),
        ("Utterance",      "utterance",        45),
        ("Response",       "response_text",    55),
        ("Path",           "path",             10),
        ("Tool Exp",       "tool_expected",    9),
        ("Tool Called",    "tool_called",      10),
        ("Tool Match",     "tool_match",       9),
        ("Follow-up",      "tool_followup",    10),
        ("Exec Tools",     "executed_tools",   20),
        ("Acc Score",      "accuracy_score",   10),
        ("Judge Reason",   "judge_reason",     40),
        ("wall_ms",        "wall_ms",          9),
        ("planner_ms",     "planner_ms",       10),
        ("retrieval_ms",   "retrieval_ms",     12),
        ("tool_ms",        "tool_ms",          8),
        ("responder_ms",   "responder_ms",     12),
        ("ttif_ms",        "ttif_ms",          9),
        ("Fast Path",      "fast_path",        9),
        ("Timestamp",      "timestamp",        18),
        ("Error",          "error",            28),
    ]
    hdr(ws_st, [c[0] for c in RAW_COLS])
    for i, r in enumerate([x for x in all_rows if x.test_type == "st"], 2):
        alt = ALT_FILL if i % 2 == 0 else PatternFill()
        vals = [run_no if a is None else
                ("Yes" if isinstance(getattr(r, a, ""), bool) and getattr(r, a) else
                 "No"  if isinstance(getattr(r, a, ""), bool) else getattr(r, a, ""))
                for _, a, _ in RAW_COLS]
        ws_st.append(vals)
        acc = r.accuracy_score
        err = r.error
        for j, cell in enumerate(ws_st[i], 1):
            cell.font = CELL_FONT; cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=(j in {8,9,17,26}))
            a = RAW_COLS[j-1][1]
            if a == "accuracy_score" and isinstance(acc, float) and acc >= 0:
                cell.fill = OK_FILL if acc >= 70 else WARN_FILL if acc >= 40 else BAD_FILL
            elif a == "error" and err:
                cell.fill = BAD_FILL
            elif a in ("tool_match","fast_path","tool_followup") and cell.value == "Yes":
                cell.fill = OK_FILL
            elif a == "tool_match" and cell.value == "No":
                cell.fill = BAD_FILL
            else:
                cell.fill = alt
    for ci, (_, _, w) in enumerate(RAW_COLS, 1):
        ws_st.column_dimensions[get_column_letter(ci)].width = w
    ws_st.freeze_panes = "A2"; ws_st.auto_filter.ref = ws_st.dimensions

    # ── Sheet 6: Raw_TH ───────────────────────────────────────────────────────
    ws_th = wb.create_sheet("Raw_TH")
    TH_COLS = [
        ("Run #",          None,               7),
        ("Label",          "label",            14),
        ("Thread ID",      "thread_id",        14),
        ("Turn #",         "turn_no",          7),
        ("Lang",           "lang",             9),
        ("Archetype",      "archetype",        28),
        ("Intent",         "expected_intent",  8),
        ("Utterance",      "utterance",        45),
        ("Response",       "response_text",    55),
        ("Path",           "path",             10),
        ("Tool Exp",       "tool_expected",    9),
        ("Tool Called",    "tool_called",      10),
        ("Tool Match",     "tool_match",       9),
        ("Follow-up",      "tool_followup",    10),
        ("Exec Tools",     "executed_tools",   20),
        ("Acc Score",      "accuracy_score",   10),
        ("Judge Reason",   "judge_reason",     40),
        ("wall_ms",        "wall_ms",          9),
        ("planner_ms",     "planner_ms",       10),
        ("retrieval_ms",   "retrieval_ms",     12),
        ("ttif_ms",        "ttif_ms",          9),
        ("Fast Path",      "fast_path",        9),
        ("Timestamp",      "timestamp",        18),
        ("Error",          "error",            28),
    ]
    hdr(ws_th, [c[0] for c in TH_COLS])
    for i, r in enumerate([x for x in all_rows if x.test_type == "th"], 2):
        alt = ALT_FILL if i % 2 == 0 else PatternFill()
        vals = [run_no if a is None else
                ("Yes" if isinstance(getattr(r, a, ""), bool) and getattr(r, a) else
                 "No"  if isinstance(getattr(r, a, ""), bool) else getattr(r, a, ""))
                for _, a, _ in TH_COLS]
        ws_th.append(vals)
        acc = r.accuracy_score; err = r.error
        for j, cell in enumerate(ws_th[i], 1):
            cell.font = CELL_FONT; cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=(j in {8,9,17,24}))
            a = TH_COLS[j-1][1]
            if a == "accuracy_score" and isinstance(acc, float) and acc >= 0:
                cell.fill = OK_FILL if acc >= 70 else WARN_FILL if acc >= 40 else BAD_FILL
            elif a == "error" and err:
                cell.fill = BAD_FILL
            elif a in ("tool_match","fast_path","tool_followup") and cell.value == "Yes":
                cell.fill = OK_FILL
            elif a == "tool_match" and cell.value == "No":
                cell.fill = BAD_FILL
            else:
                cell.fill = alt
    for ci, (_, _, w) in enumerate(TH_COLS, 1):
        ws_th.column_dimensions[get_column_letter(ci)].width = w
    ws_th.freeze_panes = "A2"; ws_th.auto_filter.ref = ws_th.dimensions

    wb.save(xl_path)
    print(f"\n  Saved JSONL -> {jsonl_path}")
    print(f"  Saved Excel -> {xl_path}  (run #{run_no:03d})")
    return xl_path


# ── Main ───────────────────────────────────────────────────────────────────────
async def main(label: str, concurrency: int, dataset: Path):
    print(f"\nFull Eval  label={label}  dataset={dataset.name}  ({time.strftime('%Y-%m-%d %H:%M')})")
    st_df = pd.read_excel(dataset, sheet_name="Golden_SingleTurn")
    th_df = pd.read_excel(dataset, sheet_name="Golden_Threads")
    st_metas = st_df.to_dict("records")
    th_ids   = list(th_df["thread_id"].unique())
    print(f"ST: {len(st_metas)} rows  |  Threads: {len(th_ids)} ({len(th_df)} turns)\n")

    keeper   = TokenKeeper()
    sem      = asyncio.Semaphore(concurrency)
    all_rows: list[Row] = []

    async with httpx.AsyncClient(timeout=120) as client:
        print("-- Single-Turn --")
        tasks = [run_st(sem, client, m, keeper, i+1, len(st_metas), label)
                 for i, m in enumerate(st_metas)]
        all_rows.extend(await asyncio.gather(*tasks))

        print("\n-- Threads --")
        for t_idx, tid in enumerate(th_ids, 1):
            t_rows = (th_df[th_df["thread_id"] == tid]
                      .sort_values("turn_no").to_dict("records"))
            keeper._issued = 0
            all_rows.extend(
                await run_thread(client, t_rows, keeper, t_idx, len(th_ids), label))

    print_summary(all_rows, label)
    save_results(all_rows, label)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--label",       default="inbound_full_v2")
    ap.add_argument("--concurrency", type=int, default=1,
                    help="parallel ST workers (default 1 — avoids Cerebras rate limits)")
    ap.add_argument("--dataset",     default=None,
                    help="path to dataset xlsx (default: highland_greenz_golden_dataset.xlsx)")
    args = ap.parse_args()
    ds = Path(args.dataset) if args.dataset else DATASET
    asyncio.run(main(args.label, args.concurrency, ds))
